import json
import os
import re
import socket
import queue
import threading
import time
from datetime import datetime
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen
import sys
import traceback
import tempfile
import faulthandler

# Install diagnostics before importing Tk, local modules, or Excel dependencies.
# pythonw.exe has no console, so otherwise an import failure looks like no launch.
CRASH_LOG = None
_CRASH_STREAM = None


def report_crash(exc_type, exc_value, exc_tb):
    detail = "".join(traceback.format_exception(exc_type, exc_value, exc_tb))
    if _CRASH_STREAM is not None:
        try:
            _CRASH_STREAM.write("\n" + detail)
            _CRASH_STREAM.flush()
        except Exception:
            pass
    text = "%s: %s\n\nError log: %s\n\nKeep this program and v6_common.py in the same folder." % (
        exc_type.__name__, exc_value, CRASH_LOG or "could not create log")
    try:
        # Native dialog also works when importing tkinter itself failed.
        if os.name == "nt":
            import ctypes
            ctypes.windll.user32.MessageBoxW(None, text, "Cred Entry error", 0x10)
        elif sys.stderr is not None:
            sys.stderr.write(detail)
    except Exception:
        pass


def install_crash_logging():
    global CRASH_LOG, _CRASH_STREAM
    for directory in (Path(__file__).resolve().parent, Path(tempfile.gettempdir())):
        try:
            path = directory / "cred_entry_crash.log"
            _CRASH_STREAM = path.open("a", encoding="utf-8", buffering=1)
            CRASH_LOG = path
            _CRASH_STREAM.write("\nStarted %s\nPython: %s\nExecutable: %s\nWindows: %s\n" % (
                datetime.now().isoformat(), sys.version, sys.executable,
                sys.getwindowsversion() if os.name == "nt" else sys.platform))
            break
        except OSError:
            continue
    sys.excepthook = report_crash
    if _CRASH_STREAM is not None:
        try:
            faulthandler.enable(file=_CRASH_STREAM)
        except (OSError, RuntimeError):
            pass


if __name__ == "__main__":
    install_crash_logging()

import shutil
import sqlite3
import uuid
from contextlib import contextmanager
from decimal import Decimal, InvalidOperation
import tkinter as tk
from tkinter import messagebox, simpledialog, ttk

import openpyxl
from openpyxl.styles import PatternFill

from v6_common import BANKS, CASHIERS

# ---------------------------------------------------------------------------
# Durable local payment ledger and conservative server synchronization.
# SQLite is authoritative; Excel files are repairable exports. An interrupted
# POST is never automatically repeated because the server has no idempotency
# contract.
# ---------------------------------------------------------------------------
HEADERS = ['ID', 'Timestamp', 'Cashier', 'Bank', 'Credit', 'Status', 'ServerEntryID', 'SmsID']


class RejectedRequest(RuntimeError):
    def __init__(self, message, status):
        super().__init__(message)
        self.status = status


def valid_amount(value):
    try:
        amount = Decimal(str(value).replace(',', '').strip())
    except InvalidOperation as exc:
        raise ValueError('Enter a valid positive amount.') from exc
    if not amount.is_finite() or amount <= 0 or amount.as_tuple().exponent < -2:
        raise ValueError('Enter a positive amount with at most two decimal places.')
    return format(amount, '.2f')


class Ledger:
    def __init__(self, app_dir):
        self.root = Path(app_dir).resolve()
        self.path = self.root / 'transactions.sqlite3'
        with self.connect() as db:
            db.executescript('''
                CREATE TABLE IF NOT EXISTS sessions (
                    folder TEXT PRIMARY KEY, business_date TEXT NOT NULL,
                    version INTEGER NOT NULL DEFAULT 0, exported INTEGER NOT NULL DEFAULT 0
                );
                CREATE TABLE IF NOT EXISTS entries (
                    transaction_id TEXT PRIMARY KEY, folder TEXT NOT NULL,
                    cashier TEXT NOT NULL, local_id INTEGER NOT NULL,
                    payload TEXT NOT NULL, server_id TEXT,
                    state TEXT NOT NULL, error TEXT NOT NULL DEFAULT '',
                    attempts INTEGER NOT NULL DEFAULT 0, next_attempt REAL NOT NULL DEFAULT 0,
                    UNIQUE(folder, cashier, local_id)
                );
            ''')

    @contextmanager
    def connect(self):
        db = sqlite3.connect(self.path, timeout=15)
        try:
            db.row_factory = sqlite3.Row
            db.execute('PRAGMA synchronous=FULL')
            with db:
                yield db
        finally:
            db.close()

    def recover_interrupted(self):
        # Called once on app startup, never while a worker is active.
        with self.connect() as db:
            db.execute("UPDATE entries SET state='uncertain', error='App stopped during submission; verify on server before retrying.' WHERE state='sending'")
            db.execute("UPDATE entries SET state='reverse_uncertain', error='App stopped during reversal; verify on server before retrying.' WHERE state='reversing'")

    def import_workbooks(self):
        """Import each old session once, atomically, without changing the workbooks."""
        for path in sorted(self.root.glob('*/marked_*.xlsx')):
            folder = path.parent.name
            if path.name != 'marked_%s.xlsx' % folder:
                continue
            with self.connect() as db:
                if db.execute('SELECT 1 FROM sessions WHERE folder=?', (folder,)).fetchone():
                    continue
            match = re.search(r'\d{4}-\d{2}-\d{2}', folder)
            if not match:
                continue
            business_date = match.group()
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                records = list(wb.active.values)
            finally:
                wb.close()
            if not records or not set(HEADERS).issubset(records[0]):
                raise ValueError('Unrecognized backup headers: %s' % path)
            # Keep the pre-migration files before Excel becomes a generated export.
            archive = path.parent / 'before-ledger-migration'
            archive.mkdir(exist_ok=True)
            for original in (path, path.parent / ('clean_%s.xlsx' % folder)):
                if original.exists() and not (archive / original.name).exists():
                    shutil.copy2(original, archive / original.name)
            with self.connect() as db:
                db.execute('INSERT INTO sessions(folder,business_date) VALUES (?,?)', (folder, business_date))
                for values in records[1:]:
                    row = dict(zip(records[0], values))
                    if row.get('ID') is None:
                        continue
                    cashier = str(row.get('Cashier') or '')
                    local_id = int(row['ID'])
                    sid = row.get('ServerEntryID')
                    state = 'reversed' if row.get('Status') == 'reversed' else ('synced' if sid else 'uncertain')
                    tid = str(uuid.uuid4())
                    payload = dict(client_transaction_id=tid, local_excel_id=local_id,
                                   timestamp=str(row.get('Timestamp') or ''), session_date=business_date,
                                   cashier=cashier, bank=str(row.get('Bank') or ''),
                                   credit=str(row.get('Credit') or '0'), source_pc=folder[:match.start()].rstrip('_'),
                                   sms_payment_id=row.get('SmsID') or None)
                    error = 'Legacy entry has no server acknowledgement; verify before retrying.' if state == 'uncertain' else ''
                    db.execute('INSERT INTO entries(transaction_id,folder,cashier,local_id,payload,server_id,state,error) VALUES (?,?,?,?,?,?,?,?)',
                               (tid, folder, cashier, local_id, json.dumps(payload), str(sid) if sid else None, state, error))

    def enqueue(self, folder, payload):
        folder = Path(folder).name
        payload = dict(payload)
        payload['credit'] = valid_amount(payload['credit'])
        tid = str(uuid.uuid4())
        with self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            sms_id = payload.get('sms_payment_id')
            if sms_id is not None:
                for existing in db.execute("SELECT payload FROM entries WHERE state<>'reversed'"):
                    if str(json.loads(existing['payload']).get('sms_payment_id')) == str(sms_id):
                        raise ValueError('This SMS already has a saved entry. Resolve or reverse that entry first.')
            db.execute('INSERT OR IGNORE INTO sessions(folder,business_date) VALUES (?,?)', (folder, payload['session_date']))
            local_id = db.execute('SELECT COALESCE(MAX(local_id),0)+1 FROM entries WHERE folder=? AND cashier=?', (folder, payload['cashier'])).fetchone()[0]
            payload.update(client_transaction_id=tid, local_excel_id=local_id)
            db.execute('INSERT INTO entries(transaction_id,folder,cashier,local_id,payload,state) VALUES (?,?,?,?,?,?)',
                       (tid, folder, payload['cashier'], local_id, json.dumps(payload), 'pending'))
            db.execute('UPDATE sessions SET version=version+1 WHERE folder=?', (folder,))
        return tid

    def rows(self, folder=None, cashier=None):
        sql, params = 'SELECT * FROM entries WHERE 1=1', []
        if folder is not None:
            sql += ' AND folder=?'
            params.append(Path(folder).name)
        if cashier is not None:
            sql += ' AND cashier=?'
            params.append(cashier)
        with self.connect() as db:
            return [dict(row) for row in db.execute(sql + ' ORDER BY rowid', params)]

    def last_cashier(self, folder):
        """Cashier of the most recent entry in a session folder, or None if it has none."""
        with self.connect() as db:
            row = db.execute('SELECT cashier FROM entries WHERE folder=? ORDER BY rowid DESC LIMIT 1',
                             (Path(folder).name,)).fetchone()
        return row['cashier'] if row else None

    def change(self, tid, state, error='', server_id=None, next_attempt=0):
        with self.connect() as db:
            db.execute('UPDATE entries SET state=?,error=?,server_id=COALESCE(?,server_id),next_attempt=? WHERE transaction_id=?',
                       (state, str(error), str(server_id) if server_id is not None else None, next_attempt, tid))
            db.execute('UPDATE sessions SET version=version+1 WHERE folder=(SELECT folder FROM entries WHERE transaction_id=?)', (tid,))

    def claim(self, tid, state, sending_state):
        with self.connect() as db:
            return db.execute('UPDATE entries SET state=?,attempts=attempts+1 WHERE transaction_id=? AND state=?',
                              (sending_state, tid, state)).rowcount == 1

    def request_reverse(self, tid):
        with self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            row = db.execute('SELECT * FROM entries WHERE transaction_id=?', (tid,)).fetchone()
            if row is None:
                raise ValueError('Select a saved local entry.')
            if row['state'] in ('sending', 'uncertain', 'reversing', 'reverse_uncertain'):
                raise ValueError('This transaction needs server verification in Sync Issues before it can be reversed.')
            if row['state'] in ('reversed', 'reverse_pending'):
                return
            state = 'reverse_pending' if row['server_id'] else 'reversed'
            db.execute("UPDATE entries SET state=?,error='',next_attempt=0 WHERE transaction_id=?", (state, tid))
            db.execute('UPDATE sessions SET version=version+1 WHERE folder=?', (row['folder'],))

    def retry_reviewed(self, tid):
        with self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            row = db.execute('SELECT * FROM entries WHERE transaction_id=?', (tid,)).fetchone()
            if not row or row['state'] not in ('uncertain', 'rejected', 'reverse_uncertain', 'reverse_rejected'):
                raise ValueError('Only transactions requiring review can be retried.')
            payload = json.loads(row['payload'])
            valid_amount(payload['credit'])
            state = 'reverse_pending' if row['state'].startswith('reverse_') else 'pending'
            db.execute("UPDATE entries SET state=?,error='',next_attempt=0 WHERE transaction_id=?", (state, tid))
            db.execute('UPDATE sessions SET version=version+1 WHERE folder=?', (row['folder'],))

    def cancel_reviewed(self, tid):
        with self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            row = db.execute('SELECT * FROM entries WHERE transaction_id=?', (tid,)).fetchone()
            if not row or row['server_id'] or row['state'] not in ('pending', 'uncertain', 'rejected'):
                raise ValueError('Only an unsent or reviewed unacknowledged entry can be cancelled here.')
            db.execute("UPDATE entries SET state='reversed',error='Cancelled after user verification' WHERE transaction_id=?", (tid,))
            db.execute('UPDATE sessions SET version=version+1 WHERE folder=?', (row['folder'],))

    def reconcile(self, server_rows):
        """Only reconcile exact server IDs or transaction UUIDs, never fuzzy matches."""
        with self.connect() as db:
            for remote in server_rows:
                sid = str(remote['id'])
                tid = remote.get('client_transaction_id')
                local = db.execute('SELECT * FROM entries WHERE server_id=? OR transaction_id=?', (sid, tid)).fetchone()
                if not local or remote.get('cashier') != local['cashier']:
                    continue
                state = local['state']
                if remote.get('status') == 'reversed':
                    new_state = 'reversed'
                elif state in ('reverse_pending', 'reversing', 'reverse_uncertain', 'reverse_rejected', 'reversed'):
                    continue  # An active snapshot must not undo an intended reversal.
                else:
                    new_state = 'synced'
                if state != new_state or local['server_id'] != sid:
                    db.execute("UPDATE entries SET server_id=?,state=?,error='' WHERE transaction_id=?", (sid, new_state, local['transaction_id']))
                    db.execute('UPDATE sessions SET version=version+1 WHERE folder=?', (local['folder'],))

    def export_dirty(self):
        errors = []
        with self.connect() as db:
            sessions = [dict(r) for r in db.execute('SELECT * FROM sessions WHERE version<>exported')]
        for session in sessions:
            folder = session['folder']
            directory = self.root / folder
            directory.mkdir(exist_ok=True)
            rows = self.rows(folder)
            try:
                for kind in ('marked', 'clean'):
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = 'Entries'
                    ws.append(HEADERS)
                    for row in rows:
                        p = json.loads(row['payload'])
                        reversed_row = row['state'] == 'reversed'
                        if kind == 'clean' and reversed_row:
                            continue
                        ws.append([row['local_id'], p['timestamp'], row['cashier'], p['bank'], p['credit'],
                                   'reversed' if reversed_row else ('active' if row['state'] == 'synced' else row['state']),
                                   row['server_id'], p.get('sms_payment_id')])
                        if reversed_row:
                            for cell in ws[ws.max_row]:
                                cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
                    target = directory / ('%s_%s.xlsx' % (kind, folder))
                    temp = target.with_name(target.name + '.' + uuid.uuid4().hex + '.tmp')
                    try:
                        wb.save(temp)
                        os.replace(temp, target)
                    finally:
                        wb.close()
                        try:
                            temp.unlink()
                        except FileNotFoundError:
                            pass
                with self.connect() as db:
                    db.execute('UPDATE sessions SET exported=? WHERE folder=? AND version=?',
                               (session['version'], folder, session['version']))
            except Exception as exc:
                errors.append('%s: %s' % (folder, exc))
        return errors


class SyncEngine:
    def __init__(self, ledger, api):
        self.ledger, self.api = ledger, api
        self.history_cursor = 0

    def sync(self):
        # Work is bounded so a large backlog does not delay UI updates indefinitely.
        candidates = [r for r in self.ledger.rows() if r['state'] in ('pending', 'reverse_pending') and r['next_attempt'] <= time.time()]
        for row in candidates[:10]:
            reverse = row['state'] == 'reverse_pending'
            sending = 'reversing' if reverse else 'sending'
            if not self.ledger.claim(row['transaction_id'], row['state'], sending):
                continue
            try:
                if reverse:
                    result = self.api.reverse_entry(row['server_id'], row['cashier'], 'cashier reversal')
                else:
                    result = self.api.create_entry(json.loads(row['payload']))
                if not isinstance(result, dict) or result.get('id') is None:
                    raise ValueError('Server response did not include an entry ID.')
                self.ledger.change(row['transaction_id'], 'reversed' if reverse else 'synced', server_id=result['id'])
            except RejectedRequest as exc:
                if exc.status == 429:
                    delay = min(300, 5 * 2 ** min(row['attempts'], 6))
                    self.ledger.change(row['transaction_id'], row['state'], str(exc), next_attempt=time.time() + delay)
                else:
                    self.ledger.change(row['transaction_id'], 'reverse_rejected' if reverse else 'rejected', str(exc))
            except Exception as exc:
                # A timeout, 5xx, or malformed reply may follow a committed write.
                self.ledger.change(row['transaction_id'], 'reverse_uncertain' if reverse else 'uncertain',
                                   'Verify on server before retrying: %s' % exc)

    def link_existing(self, tid, server_id):
        local = next((r for r in self.ledger.rows() if r['transaction_id'] == tid), None)
        if not local or local['state'] not in ('uncertain', 'rejected', 'reverse_uncertain', 'reverse_rejected'):
            raise ValueError('Only a transaction needing review can be linked.')
        payload = json.loads(local['payload'])
        rows = self.api.entries(date_from=payload['session_date'], date_to=payload['session_date'])
        remote = next((r for r in rows if str(r['id']) == str(server_id)), None)
        if remote is None:
            raise ValueError('Server entry was not found in the returned history. No change was made.')
        if (remote.get('cashier') != payload['cashier'] or remote.get('bank') != payload['bank']
                or valid_amount(remote.get('credit')) != valid_amount(payload['credit'])
                or remote.get('timestamp') != payload['timestamp']):
            raise ValueError('Server entry does not match cashier, bank, amount and timestamp. No change was made.')
        if any(r['server_id'] == str(server_id) and r['transaction_id'] != tid for r in self.ledger.rows()):
            raise ValueError('That server entry is already linked to another local transaction.')
        if local['server_id'] and local['server_id'] != str(server_id):
            raise ValueError('This transaction already has a different server ID.')
        if local['state'].startswith('reverse_') and remote.get('status') != 'reversed':
            raise ValueError('Server entry is still active. Verify the reversal before explicitly retrying it.')
        self.ledger.change(tid, 'reversed' if remote.get('status') == 'reversed' else 'synced', server_id=server_id)

    def cycle(self, session_date, review_actions=()):
        result = {'online': False, 'errors': [], 'sms': None, 'entries': None, 'config': None}
        try:
            result['config'] = self.api.config()
            result['online'] = True
        except Exception as exc:
            result['errors'].append(str(exc))
        if result['online']:
            for tid, sid in review_actions:
                try:
                    self.link_existing(tid, sid)
                except Exception as exc:
                    result['errors'].append('Verify entry %s: %s' % (sid, exc))
            try:
                self.sync()
            except Exception as exc:
                result['errors'].append('Sync: %s' % exc)
            try:
                result['sms'] = self.api.sms()
            except Exception as exc:
                result['errors'].append('SMS: %s' % exc)
            # Rotate through historical dates too, without letting history scans
            # monopolize a cycle. Missing results never prove an entry is absent.
            dates = {session_date}
            history = sorted({json.loads(r['payload'])['session_date'] for r in self.ledger.rows()} - dates)
            if history:
                for offset in range(min(2, len(history))):
                    dates.add(history[(self.history_cursor + offset) % len(history)])
                self.history_cursor = (self.history_cursor + 2) % len(history)
            for date in sorted(dates):
                try:
                    rows = self.api.entries(date_from=date, date_to=date)
                    self.ledger.reconcile(rows)
                    if date == session_date:
                        result['entries'] = rows
                except Exception as exc:
                    result['errors'].append('History %s: %s' % (date, exc))
        elif review_actions:
            result['errors'].append('Server verification could not run while offline. Queue the verification again when online.')
        result['errors'].extend(self.ledger.export_dirty())
        result['issues'] = sum(r['state'] not in ('synced', 'reversed') for r in self.ledger.rows())
        return result


APP_DIR = Path(__file__).resolve().parent
DATA_DIR = Path("C:/client-data/credit-entry-data")
DATA_DIR.mkdir(parents=True, exist_ok=True)
SESSION_STATE_FILE = DATA_DIR / "session_state.json"
CONFIG_FILE = DATA_DIR / "client_config.json"
CURRENT_SESSION_DIRECTORY = None
CURRENT_MARKED_FILE = None
CURRENT_CLEAN_FILE = None
CURRENT_SESSION_DATE = None  # business date = date the session started (survives midnight)


def load_config():
    if not os.path.exists(CONFIG_FILE):
        cfg = {"server_url": "http://192.168.1.153:8765"}
        save_config(cfg)
        return cfg
    with open(CONFIG_FILE, "r", encoding="utf-8") as f:
        cfg = json.load(f)
    cfg.setdefault("server_url", "http://127.0.0.1:8765")
    return cfg


def save_config(cfg):
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(cfg, f, indent=2, sort_keys=True)


class Api:
    def __init__(self, cfg):
        self.cfg = cfg

    def url(self, path, params=None):
        base = self.cfg["server_url"].rstrip("/")
        return base + path + (("?" + urlencode(params)) if params else "")

    def request(self, method, path, payload=None, params=None):
        data = None
        headers = {}
        if payload is not None:
            data = json.dumps(payload).encode("utf-8")
            headers["Content-Type"] = "application/json"
        if self.cfg.get("client_token"):
            headers["X-Cred-Token"] = self.cfg.get("client_token", "")
        req = Request(self.url(path, params), data=data, headers=headers, method=method)
        try:
            with urlopen(req, timeout=8) as resp:
                return json.loads(resp.read().decode("utf-8"))
        except HTTPError as exc:
            try:
                err = json.loads(exc.read().decode("utf-8")).get("error")
            except Exception:
                err = str(exc)
            finally:
                exc.close()
            if 400 <= exc.code < 500 and exc.code != 408:
                raise RejectedRequest(err or str(exc), exc.code) from exc
            raise RuntimeError(err or str(exc)) from exc
        except URLError as exc:
            raise RuntimeError("Cannot reach server: %s" % exc.reason)

    def config(self):
        return self.request("GET", "/api/config")

    def sms(self):
        return self.request("GET", "/api/sms", params={"limit": "300"})["sms"]

    def entries(self, date_from=None, date_to=None):
        params = {"limit": "500"}
        if date_from:
            params["date_from"] = date_from
        if date_to:
            params["date_to"] = date_to
        return self.request("GET", "/api/entries", params=params)["entries"]

    def create_entry(self, payload):
        return self.request("POST", "/api/entries", payload)["entry"]

    def reverse_entry(self, entry_id, cashier, reason):
        return self.request("POST", "/api/entries/reverse", {"entry_id": entry_id, "cashier": cashier, "reason": reason})["entry"]

    def add_cashier(self, name):
        return self.request("POST", "/api/cashiers", {"name": name})["cashiers"]


def today():
    return datetime.now().strftime("%Y-%m-%d")


def _session_date_from_dir(session_dir: str) -> str:
    # Folder is named "<host>_<YYYY-MM-DD>" (optionally "_<n>"); pull the date.
    m = re.search(r"(\d{4}-\d{2}-\d{2})", Path(session_dir).name)
    return m.group(1) if m else today()


def read_session_state():
    """The saved session state, or None when the previous session was ended."""
    if not Path(SESSION_STATE_FILE).exists():
        return None
    try:
        return json.loads(Path(SESSION_STATE_FILE).read_text())
    except (OSError, ValueError) as exc:
        raise RuntimeError("Cannot restore session state: %s" % exc) from exc


def write_session_state(state):
    state_path = Path(SESSION_STATE_FILE)
    temp_path = state_path.with_suffix(".tmp")
    temp_path.write_text(json.dumps(state, indent=2))
    os.replace(temp_path, state_path)


def get_session_files(cashier=None):
    global CURRENT_SESSION_DIRECTORY, CURRENT_MARKED_FILE, CURRENT_CLEAN_FILE, CURRENT_SESSION_DATE
    if CURRENT_SESSION_DIRECTORY and Path(CURRENT_SESSION_DIRECTORY).exists():
        return
    if Path(SESSION_STATE_FILE).exists():
        try:
            state = json.loads(Path(SESSION_STATE_FILE).read_text())
            saved_dir = state.get("session_directory")
            session_dir = DATA_DIR / Path(saved_dir).name if saved_dir else None
            if session_dir and Path(session_dir).exists():
                CURRENT_SESSION_DIRECTORY = str(session_dir)
                # Prefer the stored session_date; fall back to the folder name.
                CURRENT_SESSION_DATE = state.get("session_date") or _session_date_from_dir(session_dir)
                name = Path(session_dir).name
                CURRENT_MARKED_FILE = Path(session_dir) / ("marked_%s.xlsx" % name)
                CURRENT_CLEAN_FILE = Path(session_dir) / ("clean_%s.xlsx" % name)
                return
        except (OSError, ValueError, TypeError) as exc:
            raise RuntimeError("Cannot restore session state: %s" % exc) from exc
        raise RuntimeError("Saved session folder is missing from the data directory. Restore it before continuing.")
    session_date = today()  # business date fixed at session start
    base = "%s_%s" % (socket.gethostname(), session_date)
    session_dir = DATA_DIR / base
    i = 1
    while Path(session_dir).exists():
        session_dir = DATA_DIR / ("%s_%s" % (base, i))
        i += 1
    Path(session_dir).mkdir(exist_ok=True)
    CURRENT_SESSION_DIRECTORY = str(session_dir)
    CURRENT_SESSION_DATE = session_date
    name = Path(session_dir).name
    CURRENT_MARKED_FILE = Path(session_dir) / ("marked_%s.xlsx" % name)
    CURRENT_CLEAN_FILE = Path(session_dir) / ("clean_%s.xlsx" % name)
    write_session_state({"session_directory": str(session_dir), "session_date": session_date, "cashier": cashier})


def current_session_date() -> str:
    """Business date of the active session (its start date). Ensures a session exists."""
    get_session_files()
    return CURRENT_SESSION_DATE or today()


def end_current_session():
    global CURRENT_SESSION_DIRECTORY, CURRENT_MARKED_FILE, CURRENT_CLEAN_FILE, CURRENT_SESSION_DATE
    if Path(SESSION_STATE_FILE).exists():
        Path(SESSION_STATE_FILE).unlink()
    CURRENT_SESSION_DIRECTORY = None
    CURRENT_SESSION_DATE = None
    CURRENT_MARKED_FILE = None
    CURRENT_CLEAN_FILE = None


class CashierFrame(ttk.Frame):
    COLUMNS = 3

    def __init__(self, master, api, app_cfg):
        ttk.Frame.__init__(self, master, padding=24)
        self.api = api
        self.app_cfg = app_cfg
        self.pack(fill="both", expand=True)
        ttk.Style().configure("Cashier.TButton", font=("Segoe UI", 11), padding=(10, 8))
        ttk.Label(self, text="Select Cashier", font=("Segoe UI", 16, "bold")).pack(pady=(0, 4))
        owner = master.session_owner()
        if owner:
            ttk.Label(self, text="Session in progress: %s" % owner, foreground="#b45309",
                      font=("Segoe UI", 10, "bold")).pack(pady=(0, 8))
        # Fixed-width buttons in a centered grid instead of full-width bars.
        grid = ttk.Frame(self)
        grid.pack(pady=8)
        for idx, cashier in enumerate(app_cfg.get("cashiers", CASHIERS)):
            ttk.Button(grid, text=cashier, width=18, style="Cashier.TButton",
                       command=lambda c=cashier: master.open_main(api, app_cfg, c)).grid(
                row=idx // self.COLUMNS, column=idx % self.COLUMNS, padx=6, pady=6)
        ttk.Button(self, text="+ Add Cashier", width=18, command=self.add_cashier).pack(pady=(16, 0))

    def add_cashier(self):
        name = simpledialog.askstring("Add Cashier", "New cashier name:", parent=self)
        if name is None:
            return
        name = name.strip()
        if not name:
            return
        try:
            cashiers = self.api.add_cashier(name)
        except Exception as exc:
            messagebox.showerror("Add Cashier", str(exc))
            return
        self.app_cfg["cashiers"] = cashiers
        # Rebuild the selection screen so the new cashier appears.
        self.master.open_cashiers(self.api, self.app_cfg)


class MainFrame(ttk.Frame):
    def __init__(self, master, api, app_cfg, cashier):
        ttk.Frame.__init__(self, master, padding=8)
        self.api = api
        self.app_cfg = app_cfg
        self.cashier = cashier
        self.selected_sms = None
        self.selected_entry_id = None
        self.bank_var = tk.StringVar()
        self.credit_var = tk.StringVar()
        self.sms_all = []          # SMS rows in the order returned by the server
        self.sms_page = 0          # current page index for the SMS list
        self.SMS_PAGE_SIZE = 25
        self.entries = []
        self.ledger = master.ledger
        self.engine = SyncEngine(self.ledger, self.api)
        self.sync_results = queue.Queue()
        self.sync_running = False
        self.next_sync_due = 0
        self.review_actions = []
        self.is_online = True
        self.server_entries_cache = []
        self.pack(fill="both", expand=True)
        self.build()
        self.refresh_entries()
        self.network_loop()

    def build(self):
        style = ttk.Style()
        style.configure("Treeview", rowheight=24)
        style.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"))
        self.columnconfigure(1, weight=1)
        self.rowconfigure(0, weight=1)
        form = ttk.LabelFrame(self, text="New Entry", padding=8)
        form.grid(row=0, column=0, sticky="ns", padx=(0, 8))
        form.columnconfigure(0, weight=1)
        form.columnconfigure(1, weight=1)
        ttk.Label(form, text="Cashier: %s" % self.cashier, font=("Segoe UI", 11, "bold")).grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 8))
        self.bank_buttons = {}
        self.bank_frame = ttk.Frame(form)
        self.bank_frame.grid(row=1, column=0, columnspan=2, sticky="ew")
        self.render_banks(self.app_cfg.get("banks", BANKS))
        row = 2
        ttk.Label(form, text="Credit Amount").grid(row=row, column=0, columnspan=2, sticky="w", pady=(10, 0))
        self.credit_entry = ttk.Entry(form, textvariable=self.credit_var)
        self.credit_entry.grid(row=row + 1, column=0, columnspan=2, sticky="ew", pady=4)
        self.credit_entry.bind("<Return>", lambda _event: self.submit())
        ttk.Button(form, text="Submit", command=self.submit).grid(row=row + 2, column=0, columnspan=2, sticky="ew", pady=4)
        tk.Button(form, text="End Session & Close", command=self.end_session_with_confirmation, bg="#b91c1c", fg="white", activebackground="#991b1b", activeforeground="white", relief="raised").grid(row=row + 3, column=0, columnspan=2, sticky="ew", pady=(16, 0))
        self.feedback = ttk.Label(form, text="")
        self.feedback.grid(row=row + 4, column=0, columnspan=2, sticky="w", pady=(8, 0))

        work_area = ttk.Frame(self)
        work_area.grid(row=0, column=1, sticky="nsew")
        work_area.columnconfigure(0, weight=1)
        work_area.rowconfigure(0, weight=1)
        work_area.rowconfigure(1, weight=1)

        log_box = ttk.LabelFrame(work_area, text="Entries Log", padding=8)
        log_box.grid(row=0, column=0, sticky="nsew", pady=(0, 8))
        log_box.rowconfigure(1, weight=1)
        log_box.columnconfigure(0, weight=1)
        self.session_label = ttk.Label(log_box, text="")
        self.session_label.grid(row=0, column=0, sticky="w")
        cols = ("local", "server", "time", "cashier", "bank", "credit", "sms", "status")
        self.entry_tree = ttk.Treeview(log_box, columns=cols, show="headings")
        entry_columns = (
            ("local", "Local", 55, "center"),
            ("server", "Server", 65, "center"),
            ("time", "Time", 145, "w"),
            ("cashier", "Cashier", 95, "w"),
            ("bank", "Bank", 120, "w"),
            ("credit", "Credit", 95, "e"),
            ("sms", "SMS", 55, "center"),
            ("status", "Status", 75, "center"),
        )
        for col, label, width, anchor in entry_columns:
            self.entry_tree.heading(col, text=label)
            self.entry_tree.column(col, width=width, minwidth=width, anchor=anchor, stretch=(col == "time"))
        self.entry_tree.grid(row=1, column=0, sticky="nsew")
        entry_y = ttk.Scrollbar(log_box, orient=tk.VERTICAL, command=self.entry_tree.yview)
        entry_x = ttk.Scrollbar(log_box, orient=tk.HORIZONTAL, command=self.entry_tree.xview)
        self.entry_tree.configure(yscrollcommand=entry_y.set, xscrollcommand=entry_x.set)
        entry_y.grid(row=1, column=1, sticky="ns")
        entry_x.grid(row=2, column=0, sticky="ew")
        self.entry_tree.bind("<<TreeviewSelect>>", self.pick_entry)
        controls = ttk.Frame(log_box)
        controls.grid(row=3, column=0, sticky="w", pady=5)
        ttk.Button(controls, text="Reverse Selected", command=self.reverse_selected).pack(side="left")
        ttk.Button(controls, text="Sync Issues", command=self.show_sync_issues).pack(side="left", padx=8)
        self.sync_status = ttk.Label(log_box, text="Checking sync...")
        self.sync_status.grid(row=4, column=0, sticky="w")
        self.session_total_label = ttk.Label(log_box, text="Total: 0.00", font=("Segoe UI", 11, "bold"))
        self.session_total_label.grid(row=3, column=0, sticky="e", pady=5)

        sms_box = ttk.LabelFrame(work_area, text="Live Incoming Payments", padding=8)
        sms_box.grid(row=1, column=0, sticky="nsew")
        sms_box.rowconfigure(1, weight=1)
        sms_box.columnconfigure(0, weight=1)
        ttk.Button(sms_box, text="Refresh", command=self.refresh_sms).grid(row=0, column=0, sticky="e", pady=(0, 4))
        self.sms_tree = ttk.Treeview(sms_box, columns=("id", "received", "status", "bank", "amount", "payer", "logged"), show="headings", height=18)
        sms_columns = (
            ("id", "ID", 55, "center"),
            ("received", "Received", 150, "w"),
            ("status", "Status", 80, "center"),
            ("bank", "Bank", 125, "w"),
            ("amount", "Amount", 100, "e"),
            ("payer", "Payer", 230, "w"),
            ("logged", "Logged By", 105, "w"),
        )
        for col, label, width, anchor in sms_columns:
            self.sms_tree.heading(col, text=label)
            self.sms_tree.column(col, width=width, minwidth=width, anchor=anchor, stretch=(col == "payer"))
        # Sharper alternating-row contrast (clear blue vs white).
        self.entry_tree.tag_configure("odd", background="#bcd4f0")
        self.entry_tree.tag_configure("even", background="#ffffff")
        self.entry_tree.tag_configure("reversed", foreground="#7f1d1d", background="#fecaca")
        self.sms_tree.tag_configure("odd", background="#bcd4f0")
        self.sms_tree.tag_configure("even", background="#ffffff")
        # Logged SMS: obvious bright-green highlight instead of a faint grey.
        self.sms_tree.tag_configure("logged", foreground="#14532d", background="#86efac")
        self.sms_tree.tag_configure("reversed", foreground="#7f1d1d", background="#fecaca")
        self.sms_tree.grid(row=1, column=0, sticky="nsew")
        sms_y = ttk.Scrollbar(sms_box, orient=tk.VERTICAL, command=self.sms_tree.yview)
        sms_x = ttk.Scrollbar(sms_box, orient=tk.HORIZONTAL, command=self.sms_tree.xview)
        self.sms_tree.configure(yscrollcommand=sms_y.set, xscrollcommand=sms_x.set)
        sms_y.grid(row=1, column=1, sticky="ns")
        sms_x.grid(row=2, column=0, sticky="ew")
        self.sms_tree.bind("<<TreeviewSelect>>", self.pick_sms)

        pager = ttk.Frame(sms_box)
        pager.grid(row=3, column=0, sticky="ew", pady=(4, 0))
        self.sms_prev_btn = ttk.Button(pager, text="< Prev", command=self.sms_prev_page)
        self.sms_prev_btn.pack(side="left")
        self.sms_page_label = ttk.Label(pager, text="")
        self.sms_page_label.pack(side="left", padx=8)
        self.sms_next_btn = ttk.Button(pager, text="Next >", command=self.sms_next_page)
        self.sms_next_btn.pack(side="left")
        
        self.offline_banner = tk.Label(self.sms_tree, text="OFFLINE - Enter manually", font=("Segoe UI", 16, "bold"), fg="white", bg="#b91c1c")

    def render_banks(self, banks):
        """(Re)build the bank buttons; the admin can change the list while a cashier is working."""
        banks = list(banks)
        for widget in self.bank_frame.winfo_children():
            widget.destroy()
        self.bank_buttons = {}
        self.app_cfg["banks"] = banks
        for idx, bank in enumerate(banks):
            btn = tk.Button(self.bank_frame, text=bank, width=15, bg="#f0f0f0", command=lambda b=bank: self.select_bank(b))
            btn.grid(row=idx // 2, column=idx % 2, padx=3, pady=3)
            self.bank_buttons[bank] = btn
        if self.bank_var.get():
            self.select_bank(self.bank_var.get())

    def select_bank(self, bank):
        self.bank_var.set(bank)
        for name, btn in self.bank_buttons.items():
            btn.config(relief="sunken" if name == bank else "raised", bg="#4db6ac" if name == bank else "#f0f0f0")

    def network_loop(self):
        try:
            try:
                result = self.sync_results.get_nowait()
            except queue.Empty:
                result = None
            if result is not None:
                self.sync_running = False
                self.next_sync_due = time.monotonic() + 3
                self.is_online = result["online"]
                if self.is_online:
                    self.offline_banner.place_forget()
                else:
                    self.offline_banner.place(relx=0, rely=0, relwidth=1, relheight=1)
                if result.get("sms") is not None:
                    self.sms_all = list(result["sms"])
                    self.refresh_sms()
                if result.get("entries") is not None:
                    self.server_entries_cache = result["entries"]
                server_cfg = result.get("config") or {}
                if server_cfg.get("banks") and server_cfg["banks"] != self.app_cfg.get("banks"):
                    self.render_banks(server_cfg["banks"])
                if server_cfg:
                    self.master.remember_server_lists(server_cfg)
                self.refresh_entries()
                self.sync_status.config(text="%s | %s pending or needing review%s" % (
                    "Online" if self.is_online else "Offline", result.get("issues", "?"),
                    " | " + "; ".join(result["errors"]) if result["errors"] else ""))
            if not self.sync_running and time.monotonic() >= self.next_sync_due:
                date = current_session_date()
                actions, self.review_actions = self.review_actions, []
                def worker():
                    try:
                        output = self.engine.cycle(date, actions)
                    except Exception as exc:
                        output = {"online": self.is_online, "errors": [str(exc)], "issues": "?"}
                    self.sync_results.put(output)
                self.sync_running = True
                try:
                    threading.Thread(target=worker, daemon=True).start()
                except Exception:
                    self.sync_running = False
                    self.review_actions = actions + self.review_actions
                    raise
        except Exception as exc:
            self.feedback.config(text="Refresh failed: %s" % exc, foreground="red")
        finally:
            self.after(1000 if self.sync_running else 3000, self.network_loop)

    def show_sync_issues(self):
        dialog = tk.Toplevel(self)
        dialog.title("Pending transactions and sync issues")
        dialog.geometry("1050x450")
        ttk.Label(dialog, text="All sessions and cashiers. Uncertain requests require server verification before retrying.").pack(anchor="w", padx=10, pady=8)
        tree = ttk.Treeview(dialog, columns=("session", "cashier", "local", "amount", "state", "error"), show="headings")
        for name in tree["columns"]:
            tree.heading(name, text=name.title())
            tree.column(name, width=120 if name != "error" else 350)
        tree.pack(fill="both", expand=True, padx=10)
        detail = ttk.Label(dialog, text="", wraplength=1000)
        detail.pack(fill="x", padx=10, pady=5)
        def reload():
            for item in tree.get_children():
                tree.delete(item)
            for row in self.ledger.rows():
                if row["state"] in ("synced", "reversed"):
                    continue
                payload = json.loads(row["payload"])
                tree.insert("", "end", iid=row["transaction_id"], values=(row["folder"], row["cashier"], row["local_id"], payload["credit"], row["state"], row["error"]))
        def retry():
            selected = tree.selection()
            if not selected:
                return
            if not messagebox.askyesno("Verify before retrying", "Have you checked the server and confirmed this operation did NOT succeed? Retrying an already accepted payment could create a duplicate.", parent=dialog):
                return
            try:
                self.ledger.retry_reviewed(selected[0])
                reload()
            except Exception as exc:
                messagebox.showerror("Retry", str(exc), parent=dialog)
        def link():
            selected = tree.selection()
            if not selected:
                return
            sid = simpledialog.askstring("Existing server entry", "Enter the verified server entry ID. The server record must match this local payment.", parent=dialog)
            if sid:
                self.review_actions.append((selected[0], sid.strip()))
                detail.config(text="Verification queued. Use Refresh after the next sync; any failure appears in the main sync status.")
        def cancel():
            selected = tree.selection()
            if not selected:
                return
            if not messagebox.askyesno("Verify before cancelling", "Confirm that this payment was NOT accepted by the server and should be cancelled locally. If it exists on the server, link it first and reverse it instead.", parent=dialog):
                return
            try:
                self.ledger.cancel_reviewed(selected[0])
                self.refresh_entries()
                reload()
            except Exception as exc:
                messagebox.showerror("Cancel entry", str(exc), parent=dialog)
        tree.bind("<<TreeviewSelect>>", lambda event: detail.config(text=str(tree.item(tree.selection()[0], "values")[-1])) if tree.selection() else None)
        buttons = ttk.Frame(dialog)
        buttons.pack(fill="x", padx=10, pady=10)
        ttk.Button(buttons, text="Refresh", command=reload).pack(side="left")
        ttk.Button(buttons, text="Retry after verification", command=retry).pack(side="left", padx=8)
        ttk.Button(buttons, text="Link existing server entry", command=link).pack(side="left")
        ttk.Button(buttons, text="Cancel after verification", command=cancel).pack(side="left", padx=8)
        reload()

    def sms_prev_page(self):
        if self.sms_page > 0:
            self.sms_page -= 1
            self._render_sms_page()

    def sms_next_page(self):
        max_page = max(0, (len(self.sms_all) - 1) // self.SMS_PAGE_SIZE)
        if self.sms_page < max_page:
            self.sms_page += 1
            self._render_sms_page()

    def refresh_sms(self):
        # Network work runs exclusively in the sync worker.
        max_page = max(0, (len(self.sms_all) - 1) // self.SMS_PAGE_SIZE)
        self.sms_page = min(self.sms_page, max_page)
        self._render_sms_page()

    def _render_sms_page(self):
        selected_id = str(self.selected_sms["id"]) if self.selected_sms else ""
        for item in self.sms_tree.get_children():
            self.sms_tree.delete(item)
        self.sms_rows = {}
        total = len(self.sms_all)
        max_page = max(0, (total - 1) // self.SMS_PAGE_SIZE) if total else 0
        start = self.sms_page * self.SMS_PAGE_SIZE
        page_rows = self.sms_all[start:start + self.SMS_PAGE_SIZE]
        for index, row in enumerate(page_rows):
            logged = ""
            if row["status"] == "logged":
                logged = row.get("logged_by") or ""
            elif row["status"] == "reversed":
                logged = "reversed by %s" % (row.get("logged_by") or "")
            self.sms_rows[str(row["id"])] = row
            tags = ["even" if index % 2 == 0 else "odd"]
            if row["status"] == "logged":
                tags.append("logged")
            elif row["status"] == "reversed":
                tags.append("reversed")
            self.sms_tree.insert("", "end", iid=str(row["id"]), values=(row["id"], row.get("received_at") or "", row["status"], row["channel"], format(row["amount"], ",.2f"), row.get("payer") or "", logged), tags=tuple(tags))
        shown_from = start + 1 if page_rows else 0
        shown_to = start + len(page_rows)
        self.sms_page_label.config(text="Page %d/%d  (%d-%d of %d)" % (
            self.sms_page + 1, max_page + 1, shown_from, shown_to, total))
        self.sms_prev_btn.config(state=("normal" if self.sms_page > 0 else "disabled"))
        self.sms_next_btn.config(state=("normal" if self.sms_page < max_page else "disabled"))
        if selected_id and selected_id in self.sms_tree.get_children():
            self.sms_tree.selection_set(selected_id)
            self.sms_tree.focus(selected_id)

    def refresh_entries(self):
        get_session_files()
        self.session_label.config(text="Session Folder: %s" % CURRENT_SESSION_DIRECTORY)
        rows = self.ledger.rows(CURRENT_SESSION_DIRECTORY, self.cashier)
        selected_id = self.selected_entry_id
        for item in self.entry_tree.get_children():
            self.entry_tree.delete(item)
        total_credit = 0.0
        for index, row in enumerate(rows):
            p = json.loads(row["payload"])
            try:
                credit = float(valid_amount(p["credit"]))
            except ValueError:
                credit = 0.0
            if row["state"] != "reversed":
                total_credit += credit
            tags = ["even" if index % 2 == 0 else "odd"]
            if row["state"] == "reversed":
                tags.append("reversed")
            self.entry_tree.insert("", "end", iid=row["transaction_id"], values=(row["local_id"], row["server_id"] or "", p["timestamp"], row["cashier"], p["bank"], format(credit, ",.2f"), p.get("sms_payment_id") or "", row["state"]), tags=tuple(tags))
        self.session_total_label.config(text="Recorded total: %s" % format(total_credit, ",.2f"))
        if selected_id and selected_id in self.entry_tree.get_children():
            self.entry_tree.selection_set(selected_id)
            self.entry_tree.focus(selected_id)

    def pick_sms(self, _event=None):
        selected = self.sms_tree.selection()
        if not selected:
            return
        row = self.sms_rows.get(selected[0])
        if not row:
            return
        if row["status"] == "logged":
            self.sms_tree.selection_remove(selected[0])
            if self.selected_sms and str(self.selected_sms["id"]) in self.sms_tree.get_children():
                self.sms_tree.selection_set(str(self.selected_sms["id"]))
                self.sms_tree.focus(str(self.selected_sms["id"]))
            self.feedback.config(text="This SMS is already logged by %s." % (row.get("logged_by") or "another cashier"), foreground="blue")
            return
        self.selected_sms = row
        self.select_bank(row["channel"])
        self.credit_var.set("%.2f" % row["amount"])
        if row["status"] == "reversed":
            self.feedback.config(text="This SMS was reversed and can be logged again after review.", foreground="blue")
        else:
            self.feedback.config(text="Selected SMS %s. Review and submit." % row["id"], foreground="green")

    def pick_entry(self, _event=None):
        selected = self.entry_tree.selection()
        self.selected_entry_id = selected[0] if selected else None

    def submit(self):
        bank = self.bank_var.get()
        if not bank:
            messagebox.showwarning("Missing Data", "Choose a bank and enter amount.")
            return
        try:
            credit = valid_amount(self.credit_var.get())
            payload = {
                "sms_payment_id": self.selected_sms["id"] if self.selected_sms else None,
                "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "session_date": current_session_date(), "cashier": self.cashier,
                "bank": bank, "credit": credit, "source_pc": socket.gethostname(),
            }
            self.ledger.enqueue(CURRENT_SESSION_DIRECTORY, payload)
        except Exception as exc:
            messagebox.showerror("Save Entry", str(exc))
            return
        self.selected_sms = None
        self.credit_var.set("")
        self.feedback.config(text="Saved safely. Queued for server sync and Excel export.", foreground="green")
        self.refresh_entries()

    def reverse_selected(self):
        selected = self.entry_tree.selection()
        if not selected:
            return
        if not messagebox.askyesno("Reverse Entry", "Reverse this entry? A server reversal will remain pending until acknowledged."):
            return
        try:
            self.ledger.request_reverse(selected[0])
        except Exception as exc:
            messagebox.showerror("Reverse Entry", str(exc))
            return
        self.feedback.config(text="Reversal saved. Pending server changes will sync automatically.", foreground="blue")
        self.refresh_entries()

    def end_session_with_confirmation(self):
        get_session_files()
        if not CURRENT_SESSION_DIRECTORY:
            messagebox.showinfo("Info", "No active session to end.")
            return
        if not messagebox.askyesno("Confirm", "End this session? Pending transactions remain saved and will resume syncing when the app is opened again."):
            return
        end_current_session()
        self.master.destroy()


class App(tk.Tk):
    def report_callback_exception(self, exc_type, exc_value, exc_tb):
        report_crash(exc_type, exc_value, exc_tb)

    def __init__(self):
        tk.Tk.__init__(self)
        self.title("Cred Entry v6")
        self.geometry("1280x760")
        try:
            self.state("zoomed")  # open maximized on Windows
        except tk.TclError:
            pass
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        # Work around the Tk 8.6.9/8.6.10 regression (bundled with Python 3.8 on
        # Windows 7) where ttk.Treeview tag background/foreground are ignored:
        # drop the default ('!disabled','!selected') entries from the style map
        # so per-row tag colors are honored again. Harmless on newer Tk.
        def _fixed_map(option):
            return [e for e in style.map("Treeview", query_opt=option)
                    if e[:2] != ("!disabled", "!selected")]
        try:
            style.map(
                "Treeview",
                foreground=_fixed_map("foreground"),
                background=_fixed_map("background"),
            )
        except tk.TclError:
            pass
            
        try:
            self.ledger = Ledger(DATA_DIR)
            self.ledger.recover_interrupted()
            self.ledger.import_workbooks()
        except Exception as exc:
            report_crash(*sys.exc_info())
            self.destroy()
            return
        self.cfg = load_config()
        self.api = Api(self.cfg)
        self.after(100, self.auto_connect)

    def clear(self):
        for widget in self.winfo_children():
            widget.destroy()

    def auto_connect(self):
        try:
            app_cfg = self.api.config()
            self.remember_server_lists(app_cfg)
        except Exception as exc:
            print("Offline mode:", exc)
            # Last lists the server sent, so admin-managed banks survive an offline start.
            app_cfg = {"cashiers": self.cfg.get("cached_cashiers") or CASHIERS,
                       "banks": self.cfg.get("cached_banks") or BANKS}
        self.open_cashiers(self.api, app_cfg)

    def remember_server_lists(self, app_cfg):
        cashiers, banks = app_cfg.get("cashiers"), app_cfg.get("banks")
        if cashiers and banks and (cashiers, banks) != (self.cfg.get("cached_cashiers"), self.cfg.get("cached_banks")):
            self.cfg["cached_cashiers"], self.cfg["cached_banks"] = cashiers, banks
            try:
                save_config(self.cfg)
            except OSError:
                pass

    def session_owner(self):
        """Cashier holding the open session on this PC, or None if it was ended."""
        try:
            state = read_session_state()
            if not state:
                return None
            # Sessions saved before the owner was recorded: use whoever logged last in it.
            return state.get("cashier") or self.ledger.last_cashier(state.get("session_directory") or "")
        except Exception:
            return None

    def open_cashiers(self, api, app_cfg):
        self.clear()
        CashierFrame(self, api, app_cfg)

    def open_main(self, api, app_cfg, cashier):
        owner = self.session_owner()
        if owner and owner.casefold() != cashier.casefold():
            # The previous cashier did not click End Session; the next one must close it explicitly.
            if not messagebox.askyesno(
                    "Session Active",
                    "There is a current session currently active by %s.\n\n"
                    "Would you like to end that session and start a new one?" % owner):
                return
            end_current_session()
        try:
            get_session_files(cashier)
            state = read_session_state()
            if state is not None and not state.get("cashier"):
                state["cashier"] = cashier
                write_session_state(state)
            self.ledger.rows(CURRENT_SESSION_DIRECTORY, cashier)
        except Exception as exc:
            messagebox.showerror(
                "Session Recovery Error",
                "Cannot open the saved session in:\n%s\n\n%s\n\n"
                "Restore the session folder here if it was stored elsewhere, "
                "and check that the Excel backup is readable." % (DATA_DIR, exc),
            )
            return
        self.clear()
        MainFrame(self, api, app_cfg, cashier)


_SINGLE_INSTANCE_HANDLE = None


def acquire_single_instance(name="CredEntryCashier", port=61999):
    """True if no other copy is running. On Windows a named mutex is used: the OS
    frees it when the process exits (even on a crash) and, unlike a loopback
    port, it cannot be blocked by Windows reserving a port range."""
    global _SINGLE_INSTANCE_HANDLE
    if os.name == "nt":
        import ctypes
        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel32.CreateMutexW.restype = ctypes.c_void_p
        handle = kernel32.CreateMutexW(None, False, "Local\\" + name)
        if handle and ctypes.get_last_error() != 183:  # ERROR_ALREADY_EXISTS
            _SINGLE_INSTANCE_HANDLE = handle
            return True
        if handle:
            return False
        # Mutex unavailable: fall through to the port check.
    sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        sock.bind(("127.0.0.1", port))
    except OSError as exc:
        # Only "address in use" means another copy holds the port.
        return getattr(exc, "winerror", None) != 10048 and exc.errno not in (98, 48, 10048)
    _SINGLE_INSTANCE_HANDLE = sock
    return True


if __name__ == "__main__":
    if not acquire_single_instance():
        warn = tk.Tk()
        warn.withdraw()
        messagebox.showwarning("Cred Entry v6", "Cred Entry is already running on this computer.")
        warn.destroy()
        raise SystemExit(0)
    App().mainloop()
